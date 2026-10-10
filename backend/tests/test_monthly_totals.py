"""Importación de totales mensuales (historial sin comprobantes)."""
from datetime import date
from tests.conftest import BASE

TODAY = date.today()
MT = f"{BASE}/monthly-totals"


def _month(n_back: int) -> tuple[int, int]:
    y, m = TODAY.year, TODAY.month - n_back
    while m <= 0:
        m += 12
        y -= 1
    return y, m


def _label(y, m):
    return f"{y}-{m:02d}"


def _ive(c, h):
    fig = c.get(f"{BASE}/analytics/income-vs-expenses", params={"months": 6, "currency": "UYU"}, headers=h).json()
    by = {}
    for trace in fig["data"]:
        for x, y in zip(trace["x"], trace["y"]):
            by.setdefault(x, {})[trace["name"]] = y
    return by


def test_upsert_list_delete(authed):
    c, h = authed
    y, m = _month(2)
    r = c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "uyu", "income": "50000", "expense": "30000"}]}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json() == {"created": 1, "updated": 0}
    # Mismo mes y moneda → se reemplaza
    r = c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "UYU", "income": "51000", "expense": "30000"}]}, headers=h)
    assert r.json() == {"created": 0, "updated": 1}
    rows = c.get(MT, headers=h).json()
    assert len(rows) == 1 and rows[0]["income"] == "51000.00" and rows[0]["currency"] == "UYU"
    assert c.delete(f"{MT}/{rows[0]['id']}", headers=h).status_code == 204
    assert c.get(MT, headers=h).json() == []


def test_future_month_rejected(authed):
    c, h = authed
    y, m = TODAY.year + 1, 1
    r = c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "UYU", "income": "1", "expense": "1"}]}, headers=h)
    assert r.status_code == 422


def test_totals_show_in_income_vs_expenses_and_patrimonio(authed, cash_account):
    c, h = authed
    y2, m2 = _month(2)
    y1, m1 = _month(1)
    c.post(MT, json={"rows": [
        {"year": y2, "month": m2, "currency": "UYU", "income": "50000", "expense": "30000"},
        {"year": y1, "month": m1, "currency": "UYU", "income": "40000", "expense": "45000"},
    ]}, headers=h)

    ive = _ive(c, h)
    assert ive[_label(y2, m2)] == {"Ingresos": 50000.0, "Egresos": 30000.0}
    assert ive[_label(y1, m1)] == {"Ingresos": 40000.0, "Egresos": 45000.0}

    pat = {p["month"]: p["value"] for p in c.get(f"{BASE}/analytics/patrimonio", params={"currency": "UYU"}, headers=h).json()}
    now = pat[TODAY.strftime("%Y-%m")]
    assert now == 10000.0  # el saldo actual no cambia
    assert pat[_label(y1, m1)] == now           # mes actual sin flujo
    assert pat[_label(*_month(2))] == now + 5000  # el mes pasado gastó 5000 más de lo que ingresó
    assert pat[_label(*_month(3))] == now + 5000 - 20000


def test_totals_in_usd_are_converted(authed):
    c, h = authed
    y, m = _month(1)
    c.post(f"{BASE}/exchange-rates", json={"from_currency": "USD", "to_currency": "UYU", "rate": "40",
                                           "year": y, "month": m}, headers=h)
    c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "USD", "income": "100", "expense": "50"}]}, headers=h)
    assert _ive(c, h)[_label(y, m)] == {"Ingresos": 4000.0, "Egresos": 2000.0}


def test_monthly_breakdown_includes_totals(authed):
    c, h = authed
    y, m = _month(1)
    c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "UYU", "income": "1000", "expense": "400"}]}, headers=h)
    b = c.get(f"{BASE}/analytics/monthly-breakdown", params={"year": y, "month": m, "currency": "UYU"}, headers=h).json()
    assert b["income"] == 1000.0 and b["expenses"] == 400.0


def test_totals_are_private(authed, client):
    c, h = authed
    y, m = _month(1)
    c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "UYU", "income": "1", "expense": "1"}]}, headers=h)
    total_id = c.get(MT, headers=h).json()[0]["id"]
    c.post(f"{BASE}/users/register", json={"name": "Otro", "email": "otro2@example.com",
                                           "password": "Test1234!", "currency_default": "UYU"})
    tok = c.post(f"{BASE}/auth/login", data={"username": "otro2@example.com", "password": "Test1234!"}).json()["access_token"]
    other = {"Authorization": f"Bearer {tok}"}
    assert c.get(MT, headers=other).json() == []
    assert c.delete(f"{MT}/{total_id}", headers=other).status_code == 404


def test_first_tx_month_includes_imported_totals(authed):
    c, h = authed
    y, m = _month(4)
    c.post(MT, json={"rows": [{"year": y, "month": m, "currency": "UYU", "income": "1", "expense": "1"}]}, headers=h)
    s = c.get(f"{BASE}/summary", headers=h).json()
    assert s["first_tx_month"] == _label(y, m)


# ─── Plantilla .xlsx / .csv ───────────────────────────────────────────────────

def _filled_template(rows) -> bytes:
    """Completa la plantilla real (frontend/public) como lo haría el usuario en Excel."""
    import io
    from pathlib import Path
    from openpyxl import load_workbook
    path = Path(__file__).resolve().parents[2] / "frontend" / "public" / "plantilla-totales-mensuales.xlsx"
    wb = load_workbook(path)
    ws = wb.active
    for i, row in enumerate(rows, start=8):  # la tabla empieza en la fila 8
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_filled_template(authed):
    from datetime import datetime
    c, h = authed
    y1, m1 = _month(2)
    y2, m2 = _month(1)
    content = _filled_template([
        [f"{y1}-{m1:02d}", "UYU", 52000, 41300.5],
        [datetime(y2, m2, 1), "usd", "1.250,00", "900"],   # Excel convirtió el mes en fecha
        [f"{TODAY.year + 1}-01", "UYU", 1, 1],               # futuro → error
        ["marzo", "UYU", 1, 1],                              # mes inválido → error
    ])
    r = c.post(f"{MT}/parse", files={"file": ("plantilla.xlsx", content)}, headers=h)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows"] == [
        {"year": y1, "month": m1, "currency": "UYU", "income": "52000.00", "expense": "41300.50"},
        {"year": y2, "month": m2, "currency": "USD", "income": "1250.00", "expense": "900.00"},
    ]
    assert [e["row"] for e in body["errors"]] == [10, 11]
    # Nada se guardó todavía
    assert c.get(MT, headers=h).json() == []


def test_parse_csv_uses_default_currency(authed):
    c, h = authed
    y, m = _month(1)
    csv_text = f"Mes;Ingresos;Gastos\n{y}-{m:02d};45.000;30.500,75\n"
    r = c.post(f"{MT}/parse", files={"file": ("totales.csv", csv_text.encode())}, data={"currency": "UYU"}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == [{"year": y, "month": m, "currency": "UYU", "income": "45000.00", "expense": "30500.75"}]


def test_parse_rejects_unknown_file(authed):
    c, h = authed
    r = c.post(f"{MT}/parse", files={"file": ("foto.png", b"\x89PNG")}, headers=h)
    assert r.status_code == 422
    r = c.post(f"{MT}/parse", files={"file": ("x.csv", b"a;b;c\n1;2;3\n")}, headers=h)
    assert r.status_code == 422
