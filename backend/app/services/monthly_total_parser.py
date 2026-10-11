"""
Lectura de la plantilla de totales mensuales (o cualquier .xlsx/.csv con columnas
Mes, Moneda, Ingresos, Gastos). No guarda nada: devuelve filas para previsualizar.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.exceptions.monthly_total_exceptions import MonthlyTotalFileError

_HEADERS = {
    "mes": "month",
    "moneda": "currency",
    "ingresos": "income",
    "ingreso": "income",
    "gastos": "expense",
    "gasto": "expense",
}


def _parse_month(value) -> tuple[int, int] | None:
    if isinstance(value, (datetime, date)):
        return value.year, value.month
    s = str(value).strip()
    m = re.match(r"^(\d{4})[-/.](\d{1,2})(?:[-/.]\d{1,2})?$", s) or None
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
    else:
        m = re.match(r"^(\d{1,2})[-/.](\d{4})$", s)
        if not m:
            return None
        y, mo = int(m.group(2)), int(m.group(1))
    return (y, mo) if 1 <= mo <= 12 else None


def _parse_amount(value) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        d = Decimal(str(value))
        return d if d >= 0 else None
    s = re.sub(r"[^\d.,-]", "", str(value))
    if not s:
        return None
    if "." in s and "," in s:
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        thou = "," if dec == "." else "."
        s = s.replace(thou, "").replace(dec, ".")
    elif "," in s or "." in s:
        sep = "," if "," in s else "."
        parts = s.split(sep)
        thousands = len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3)
        s = "".join(parts) if thousands else ".".join(parts)
    try:
        d = Decimal(s)
    except InvalidOperation:
        return None
    return d if d >= 0 else None


def _rows_from_file(content: bytes, filename: str) -> list[list]:
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        except Exception:
            raise MonthlyTotalFileError("No se pudo leer el Excel. Usá la plantilla de Fluxo (.xlsx).")
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith(".csv") or name.endswith(".txt"):
        text = content.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    raise MonthlyTotalFileError("Formato no soportado. Subí un .xlsx o .csv.")


def parse_file(content: bytes, filename: str, default_currency: str) -> dict:
    rows = _rows_from_file(content, filename)

    header_idx, cols = None, {}
    for i, row in enumerate(rows):
        found = {}
        for j, cell in enumerate(row):
            key = re.sub(r"\(.*?\)", "", str(cell or "")).strip().lower()
            if key in _HEADERS and _HEADERS[key] not in found:
                found[_HEADERS[key]] = j
        if "month" in found and ("income" in found or "expense" in found):
            header_idx, cols = i, found
            break
    if header_idx is None:
        raise MonthlyTotalFileError("No encontré las columnas Mes, Ingresos y Gastos. Usá la plantilla de Fluxo.")

    today = date.today()
    parsed, errors = [], []
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        def cell(key):
            j = cols.get(key)
            return row[j] if j is not None and j < len(row) else None

        raw_month = cell("month")
        if raw_month is None or str(raw_month).strip() == "":
            continue  # fila vacía
        month = _parse_month(raw_month)
        if not month:
            errors.append({"row": i, "message": f"Mes inválido: {raw_month} (usá AAAA-MM)"})
            continue
        if month > (today.year, today.month):
            errors.append({"row": i, "message": f"{month[0]}-{month[1]:02d} es un mes futuro"})
            continue
        income, expense = _parse_amount(cell("income")), _parse_amount(cell("expense"))
        if income is None or expense is None:
            errors.append({"row": i, "message": "Monto inválido"})
            continue
        if income == 0 and expense == 0:
            continue
        currency = str(cell("currency") or default_currency).strip().upper()
        if not re.match(r"^[A-Z]{2,10}$", currency):
            errors.append({"row": i, "message": f"Moneda inválida: {currency}"})
            continue
        parsed.append({
            "year": month[0], "month": month[1], "currency": currency,
            "income": str(income.quantize(Decimal("0.01"))),
            "expense": str(expense.quantize(Decimal("0.01"))),
        })
    return {"rows": parsed, "errors": errors}
