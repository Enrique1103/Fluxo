"""
Genera frontend/public/plantilla-totales-mensuales.xlsx: la plantilla que el usuario
descarga en Importar → Totales mensuales, completa en Excel y vuelve a subir.

Uso (desde backend/):  python scripts/generar_plantilla_totales.py
Requiere openpyxl y Pillow (Pillow solo para insertar el logo; no hace falta en el servidor).

La hoja "Mi historial" es la que lee POST /monthly-totals/parse: busca la fila de
encabezados (Mes / Moneda / Ingresos / Gastos) e ignora el resto de las columnas.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "frontend" / "public" / "pwa-512.png"  # ícono "Fluxo" de la app
OUT = ROOT / "frontend" / "public" / "plantilla-totales-mensuales.xlsx"

HEADER_ROW = 9
FIRST = HEADER_ROW + 1
DATA_ROWS = 120  # 10 años de meses
LAST = FIRST + DATA_ROWS - 1

# Colores de Fluxo (mismos que el ícono)
NIGHT = "020617"
INDIGO = "312E81"
VIOLET = "A78BFA"
LAVENDER = "F5F3FF"
SLATE = "64748B"
WHITE = "FFFFFF"
GREEN = "059669"
RED = "E11D48"

band = PatternFill("solid", fgColor=INDIGO)
night = PatternFill("solid", fgColor=NIGHT)
lavender = PatternFill("solid", fgColor=LAVENDER)
thin = Side(style="thin", color="DDD6FE")


def logo(size: int) -> XLImage:
    img = XLImage(str(LOGO))
    img.width = img.height = size
    return img


def fill_row(ws, row: int, cols: str, fill: PatternFill) -> None:
    for col in cols:
        ws[f"{col}{row}"].fill = fill


wb = Workbook()

# ─── Hoja 1: Mi historial ─────────────────────────────────────────────────────
ws = wb.active
ws.title = "Mi historial"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = VIOLET

for col, width in zip("ABCDEF", (19, 11, 17, 17, 17, 3)):
    ws.column_dimensions[col].width = width

# Banda de marca
for r in range(1, 5):
    fill_row(ws, r, "ABCDEF", band)
    ws.row_dimensions[r].height = 19
ws.add_image(logo(64), "A1")
ws.merge_cells("B1:E2")
ws["B1"] = "Mi historial en Fluxo"
ws["B1"].font = Font(name="Calibri", size=20, bold=True, color=WHITE)
ws["B1"].alignment = Alignment(vertical="bottom", indent=1)
ws.merge_cells("B3:E3")
ws["B3"] = "Tu camino financiero no empieza hoy. Traelo con vos."
ws["B3"].font = Font(name="Calibri", size=11, italic=True, color=VIOLET)
ws["B3"].alignment = Alignment(indent=1)
ws.merge_cells("B4:E4")
ws["B4"] = "Completá un mes por fila y subí este archivo en Importar → Totales mensuales."
ws["B4"].font = Font(name="Calibri", size=9, color="C7D2FE")
ws["B4"].alignment = Alignment(indent=1)

# Resumen (se calcula solo mientras completás)
labels = [("A", "Meses cargados"), ("C", "Ingresos totales"), ("D", "Gastos totales"), ("E", "Ahorro total")]
for col, text in labels:
    c = ws[f"{col}6"]
    c.value = text
    c.font = Font(size=9, bold=True, color=SLATE)
    c.alignment = Alignment(horizontal="center")
ws["A7"] = f'=COUNTA(A{FIRST}:A{LAST})'
ws["C7"] = f'=SUM(C{FIRST}:C{LAST})'
ws["D7"] = f'=SUM(D{FIRST}:D{LAST})'
ws["E7"] = f'=SUM(E{FIRST}:E{LAST})'
for col in "ACDE":
    c = ws[f"{col}7"]
    c.font = Font(size=14, bold=True, color=INDIGO)
    c.alignment = Alignment(horizontal="center")
    c.fill = lavender
    if col != "A":
        c.number_format = "#,##0"
ws["E7"].font = Font(size=14, bold=True, color=GREEN)
ws.row_dimensions[7].height = 24
ws.merge_cells("A8:E8")
ws["A8"] = "El resumen suma todas las filas: si cargás meses en distintas monedas, tomalo como referencia."
ws["A8"].font = Font(size=8, italic=True, color=SLATE)

# Encabezados de la tabla
headers = ["Mes (AAAA-MM)", "Moneda", "Ingresos", "Gastos", "Ahorro del mes"]
for col, h in enumerate(headers, start=1):
    c = ws.cell(row=HEADER_ROW, column=col, value=h)
    c.font = Font(bold=True, color=WHITE)
    c.fill = night
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = Border(bottom=Side(style="medium", color=VIOLET))
ws.row_dimensions[HEADER_ROW].height = 24

for r in range(FIRST, LAST + 1):
    for col in range(1, 6):
        c = ws.cell(row=r, column=col)
        c.border = Border(bottom=thin)
        if (r - FIRST) % 2 == 1:
            c.fill = lavender
    ws.cell(row=r, column=1).number_format = "@"  # texto: Excel no lo convierte en fecha
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3).number_format = "#,##0.00"
    ws.cell(row=r, column=4).number_format = "#,##0.00"
    e = ws.cell(row=r, column=5, value=f'=IF(AND(C{r}="",D{r}=""),"",N(C{r})-N(D{r}))')
    e.number_format = "#,##0.00;-#,##0.00"
    e.font = Font(bold=True)

ws.conditional_formatting.add(
    f"E{FIRST}:E{LAST}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], font=Font(bold=True, color=GREEN))
)
ws.conditional_formatting.add(
    f"E{FIRST}:E{LAST}", CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, color=RED))
)

dv_cur = DataValidation(type="list", formula1='"UYU,USD,EUR,ARS,BRL"', allow_blank=True)
dv_cur.errorTitle, dv_cur.error = "Moneda", "Elegí una moneda de la lista"
dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv_num.errorTitle, dv_num.error = "Monto", "El monto debe ser un número mayor o igual a 0"
dv_mes = DataValidation(type="textLength", operator="between", formula1="6", formula2="7", allow_blank=True)
dv_mes.promptTitle, dv_mes.prompt = "Mes", "Escribilo así: 2025-03"
dv_mes.showInputMessage = True
for dv, ref in ((dv_cur, f"B{FIRST}:B{LAST}"), (dv_num, f"C{FIRST}:D{LAST}"), (dv_mes, f"A{FIRST}:A{LAST}")):
    ws.add_data_validation(dv)
    dv.add(ref)

ws.freeze_panes = f"A{FIRST}"
ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.oddFooter.center.text = "Fluxo · Mi historial"

# ─── Hoja 2: Cómo completar ───────────────────────────────────────────────────
g = wb.create_sheet("Cómo completar")
g.sheet_view.showGridLines = False
g.sheet_properties.tabColor = INDIGO
for col, width in zip("ABCDEF", (4, 19, 11, 17, 17, 30)):
    g.column_dimensions[col].width = width
for r in range(1, 4):
    fill_row(g, r, "ABCDEF", band)
    g.row_dimensions[r].height = 20
g.add_image(logo(52), "B1")
g.merge_cells("C1:F2")
g["C1"] = "Cómo traer tu historial a Fluxo"
g["C1"].font = Font(size=16, bold=True, color=WHITE)
g["C1"].alignment = Alignment(vertical="center", indent=1)
g.merge_cells("C3:F3")
g["C3"] = "En 3 pasos, sin cargar cada gasto de nuevo."
g["C3"].font = Font(size=10, italic=True, color=VIOLET)
g["C3"].alignment = Alignment(indent=1)

steps = [
    ("1", "Buscá en tu app anterior el total de ingresos y de gastos de cada mes."),
    ("2", "Cargalos en la hoja \"Mi historial\": un mes por fila, mes en formato AAAA-MM."),
    ("3", "En Fluxo andá a Importar → Totales mensuales → Subir plantilla completada."),
]
r = 5
for num, text in steps:
    g[f"B{r}"] = num
    g[f"B{r}"].font = Font(size=16, bold=True, color=VIOLET)
    g[f"B{r}"].alignment = Alignment(horizontal="center", vertical="center")
    g.merge_cells(f"C{r}:F{r}")
    g[f"C{r}"] = text
    g[f"C{r}"].font = Font(size=11)
    g[f"C{r}"].alignment = Alignment(vertical="center", wrap_text=True)
    g.row_dimensions[r].height = 30
    r += 1

r += 1
g[f"B{r}"] = "Ejemplo"
g[f"B{r}"].font = Font(size=11, bold=True, color=INDIGO)
r += 1
for col, h in zip("BCDE", ["Mes (AAAA-MM)", "Moneda", "Ingresos", "Gastos"]):
    c = g[f"{col}{r}"]
    c.value, c.font, c.fill = h, Font(bold=True, color=WHITE), night
    c.alignment = Alignment(horizontal="center")
for row in (("2025-01", "UYU", 52000, 41300), ("2025-02", "UYU", 52000, 38950), ("2025-02", "USD", 300, 120)):
    r += 1
    for col, v in zip("BCDE", row):
        c = g[f"{col}{r}"]
        c.value = v
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=thin)
        if col in "DE":
            c.number_format = "#,##0"

r += 2
g[f"B{r}"] = "Bueno saber"
g[f"B{r}"].font = Font(size=11, bold=True, color=INDIGO)
tips = [
    "Antes de guardar vas a ver todos los meses en Fluxo para revisarlos.",
    "Si subís un mes que ya estaba cargado en la misma moneda, se reemplaza.",
    "Podés tener el mismo mes en UYU y en USD: Fluxo convierte con tu tasa de cambio.",
    "No se crean comprobantes: tus gráficos muestran el historial y seguís registrando día a día.",
    "Conviene crear tus cuentas con el saldo de hoy antes de importar, así el patrimonio sale bien.",
]
for tip in tips:
    r += 1
    g[f"B{r}"] = "•"
    g[f"B{r}"].font = Font(bold=True, color=VIOLET)
    g[f"B{r}"].alignment = Alignment(horizontal="right")
    g.merge_cells(f"C{r}:F{r}")
    g[f"C{r}"] = tip
    g[f"C{r}"].font = Font(size=10, color="334155")

r += 2
g.merge_cells(f"B{r}:F{r}")
g[f"B{r}"] = "Fluxo · Finanzas personales inteligentes"
g[f"B{r}"].font = Font(size=9, italic=True, color=SLATE)
g[f"B{r}"].alignment = Alignment(horizontal="center")

wb.active = 0
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Plantilla generada: {OUT}")
